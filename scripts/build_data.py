"""Convert the ESCRS IOL Calculator GA xlsx export into JSON files for the dashboard.

Run with: uv run python scripts/build_data.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "ga" / "ga17042026.xlsx"
OUT_DIR = ROOT / "data"
# Restrict to fully-covered years. 2023 is partial (Sep–Dec only) and 2026 is
# incomplete across sheets (missing from `overall`, partial in `country`/`calcs`).
KEEP_YEARS = {2024, 2025}


def _cell(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def read_sheet(ws):
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    rows = []
    for raw in rows_iter:
        if all(c is None for c in raw):
            continue
        rows.append({h: _cell(v) for h, v in zip(header, raw)})
    return rows


def dedupe(rows, keys):
    seen = set()
    out = []
    for r in rows:
        k = tuple(r.get(k) for k in keys)
        if k in seen:
            continue
        seen.add(k)
        out.append(r)
    return out


def build_overall(ws):
    rows = read_sheet(ws)
    rows = dedupe(rows, ["Year", "Month"])
    monthly = []
    yearly = []
    for r in rows:
        year = int(r["Year"])
        if year not in KEEP_YEARS:
            continue
        month = int(r["Month"])
        item = {
            "year": year,
            "month": month,
            "users": int(r["Users"] or 0),
            "sessions": int(r["Sessions"] or 0),
            "avg_session_duration": float(r["Average session duration (seconds)"] or 0.0),
        }
        # Month 0 appears to be a year-aggregate row (unique users across the year).
        if month == 0:
            yearly.append(item)
        else:
            monthly.append(item)
    monthly.sort(key=lambda r: (r["year"], r["month"]))
    yearly.sort(key=lambda r: r["year"])
    return {"monthly": monthly, "yearly": yearly}


def build_countries(ws):
    rows = read_sheet(ws)
    rows = dedupe(rows, ["Year", "Country"])
    out = []
    for r in rows:
        year = int(r["Year"])
        if year not in KEEP_YEARS:
            continue
        country = (r.get("Country") or "").strip()
        continent = (r.get("Continent") or "").strip() or "Unknown"
        out.append({
            "year": year,
            "country": country,
            "continent": continent,
            "total_users": int(r.get("Total users") or 0),
            "new_users": int(r.get("New users") or 0),
            "sessions": int(r.get("Sessions") or 0),
            "views_per_session": float(r.get("Views per session") or 0.0),
        })
    out.sort(key=lambda r: (r["year"], -r["total_users"]))
    return out


def build_calcs(ws):
    rows = read_sheet(ws)
    rows = dedupe(rows, ["Year", "Month", "Calculator", "Toric", "Keratoconus", "Post LASIK/PRK"])
    out = []
    for r in rows:
        year = int(r["Year"])
        if year not in KEEP_YEARS:
            continue
        out.append({
            "year": year,
            "month": int(r["Month"]),
            "calculator": (r.get("Calculator") or "").strip(),
            "toric": bool(int(r.get("Toric") or 0)),
            "keratoconus": bool(int(r.get("Keratoconus") or 0)),
            "post_lasik": bool(int(r.get("Post LASIK/PRK") or 0)),
            "calculations": int(r.get("Calculations") or 0),
        })
    out.sort(key=lambda r: (r["year"], r["month"], r["calculator"]))
    return out


def main():
    if not XLSX.exists():
        print(f"ERROR: {XLSX} not found", file=sys.stderr)
        sys.exit(1)
    OUT_DIR.mkdir(exist_ok=True)

    wb = load_workbook(XLSX, data_only=True)
    overall = build_overall(wb["overall"])
    countries = build_countries(wb["country"])
    calcs = build_calcs(wb["calcs"])

    meta = {
        "source": XLSX.name,
        "years": sorted({r["year"] for r in overall["monthly"]} | {r["year"] for r in countries} | {r["year"] for r in calcs}),
        "calculators": sorted({r["calculator"] for r in calcs}),
        "continents": sorted({r["continent"] for r in countries}),
        "country_count": len({r["country"] for r in countries}),
        "overall_month_count": len(overall["monthly"]),
        "calcs_row_count": len(calcs),
    }

    (OUT_DIR / "overall.json").write_text(json.dumps(overall, indent=2))
    (OUT_DIR / "countries.json").write_text(json.dumps(countries, indent=2))
    (OUT_DIR / "calcs.json").write_text(json.dumps(calcs, indent=2))
    (OUT_DIR / "meta.json").write_text(json.dumps(meta, indent=2))

    print("Wrote:")
    for p in ("overall.json", "countries.json", "calcs.json", "meta.json"):
        f = OUT_DIR / p
        print(f"  {f.relative_to(ROOT)}  ({f.stat().st_size:,} bytes)")
    print(f"\nMeta: {json.dumps(meta, indent=2)}")


if __name__ == "__main__":
    main()
